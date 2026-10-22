"""Generate binary red-team samples: PII hidden in DOCX headers/footers,
spreadsheet cells, a PDF footer, and homoglyph/zero-width obfuscation.

    .venv/Scripts/python scripts/make_samples.py
"""
from __future__ import annotations

from pathlib import Path

SAMPLES = Path(__file__).resolve().parents[1] / "samples"


def docx_with_header_pii(path: Path) -> None:
    import docx

    d = docx.Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "CONFIDENTIAL - Prepared for client Margaret Wilson, SSN 512-44-7781"
    sec.footer.paragraphs[0].text = "Account #MW-4471902 | margaret.wilson@example.org | (303) 555-0144"
    d.add_heading("Managed Portfolio Investment Agreement", 1)
    d.add_paragraph("This Investment Agreement sets out the terms of the managed portfolio service.")
    d.add_paragraph("Our managed portfolio is risk-free and we guarantee a 12% annual return.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Beneficiary", "Thomas Wilson"
    t.cell(1, 0).text, t.cell(1, 1).text = "Beneficiary SSN", "623-18-9045"
    d.add_paragraph("Investing involves risk, including the possible loss of principal.")
    d.save(path)


def xlsx_with_pii(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Client Roster"
    ws.append(["Name", "SSN", "Email", "Balance"])
    ws.append(["Kevin Murphy", "401-22-8890", "kevin.murphy@example.com", "$310,000"])
    ws.append(["Laura Nguyen", "288-91-4402", "laura.n@example.com", "$95,500"])
    notes = wb.create_sheet("Newsletter Notes")
    notes.append(["Draft copy for the quarterly newsletter"])
    notes.append(["Our flagship fund returned 31% last year and will outperform again next year."])
    notes.append(["This material is not an offer to buy or sell any security."])
    wb.save(path)


def pdf_with_footer_pii(path: Path) -> None:
    from fpdf import FPDF

    class Doc(FPDF):
        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", size=8)
            self.cell(0, 10, "Client: Angela Brooks | SSN 330-71-2219 | Acct #AB-7731005", align="C")

    pdf = Doc()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 7, "Wealth Advisory Brochure\n\n"
                         "Our advisors are completely unbiased and free of any conflicts of interest. "
                         "Zero commissions on every trade.\n\n"
                         "Investing involves risk, including the possible loss of principal. "
                         "Past performance does not guarantee future results.")
    pdf.output(str(path))


def homoglyph_txt(path: Path) -> None:
    # Cyrillic 'а'/'е'/'о', full-width digits and zero-width spaces hide PII and banned phrases.
    text = (
        "INVESTMENT AGREEMENT\n\n"
        "Client: Jаne Smith\n"  # Cyrillic а
        "SSN: ９８７-６５-４３２１\n"  # full-width digits
        "Phone: 415​-555-​0132\n\n"  # zero-width spaces
        "We guаrаntee a 15% annual return.\n\n"  # Cyrillic а in 'guarantee'
        "Investing involves risk, including the possible loss of principal.\n"
    )
    path.write_text(text, encoding="utf-8")


if __name__ == "__main__":
    SAMPLES.mkdir(exist_ok=True)
    docx_with_header_pii(SAMPLES / "RedTeam_Header_Footer_PII_Investment_Agreement.docx")
    xlsx_with_pii(SAMPLES / "RedTeam_Spreadsheet_PII_Newsletter.xlsx")
    pdf_with_footer_pii(SAMPLES / "RedTeam_PDF_Footer_PII_Brochure.pdf")
    homoglyph_txt(SAMPLES / "RedTeam_Homoglyph_Investment_Agreement.txt")
    print("samples written to", SAMPLES)
